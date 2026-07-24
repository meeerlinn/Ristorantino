"""
Genera Manual de Usuario de Ristorantino — tres secciones: Mozo, Chef, Cajero.
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ── Paleta ────────────────────────────────────────────────────────────────────
C_DARK      = "1B2B3A"
C_MID       = "2E6DA4"
C_ACCENT    = "C8A96E"
C_GRAY      = "F5F5F0"
C_TEXT      = "2C2C2C"
C_WHITE     = "FFFFFF"
C_CODE      = "F0F0EA"
C_GREEN     = "1A6B3C"  # mozo
C_ORANGE    = "B8560A"  # chef
C_PURPLE    = "5B3A8C"  # cajero
C_GREEN_L   = "E8F5EE"
C_ORANGE_L  = "FDF0E6"
C_PURPLE_L  = "F0EBF8"

def hex_to_rgb(h):
    return RGBColor(int(h[0:2],16), int(h[2:4],16), int(h[4:6],16))

# ── XML helpers ───────────────────────────────────────────────────────────────

def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def set_cell_borders(cell, color="CCCCCC"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)

def set_cell_margins(cell, top=80, bottom=80, left=140, right=140):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for side, val in (('top', top), ('bottom', bottom), ('left', left), ('right', right)):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:w'), str(val))
        el.set(qn('w:type'), 'dxa')
        tcMar.append(el)
    tcPr.append(tcMar)

def add_paragraph_border_bottom(para, hex_color, sz=8):
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(sz))
    bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), hex_color)
    pBdr.append(bottom)
    pPr.append(pBdr)

def style_run(run, size=11, bold=False, italic=False, color=C_TEXT, font='Arial'):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = hex_to_rgb(color)

# ── Componentes de contenido ──────────────────────────────────────────────────

def add_space(doc, size=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(0)
    p.add_run().font.size = Pt(size)

def add_para(doc, text, size=10.5, color=C_TEXT, bold=False, italic=False,
             align=WD_ALIGN_PARAGRAPH.LEFT, space_before=3, space_after=4):
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(space_before)
    para.paragraph_format.space_after  = Pt(space_after)
    para.alignment = align
    run = para.add_run(text)
    style_run(run, size=size, bold=bold, italic=italic, color=color)
    return para

def add_bullet(doc, text, accent_color=C_MID):
    para = doc.add_paragraph(style='List Bullet')
    para.paragraph_format.space_before = Pt(2)
    para.paragraph_format.space_after  = Pt(2)
    run = para.add_run(text)
    style_run(run, size=10, color=C_TEXT)
    return para

def add_numbered(doc, text):
    para = doc.add_paragraph(style='List Number')
    para.paragraph_format.space_before = Pt(2)
    para.paragraph_format.space_after  = Pt(3)
    run = para.add_run(text)
    style_run(run, size=10.5, color=C_TEXT)
    return para

def add_heading(doc, text, level, color=None, border_color=None):
    para = doc.add_heading(text, level=level)
    para.clear()
    sizes = {1: 20, 2: 14, 3: 12}
    sz  = sizes.get(level, 11)
    c   = color or C_DARK
    run = para.add_run(text)
    style_run(run, size=sz, bold=True, color=c)
    if border_color:
        add_paragraph_border_bottom(para, border_color, sz=10)
    return para

def add_warning_box(doc, text, bg_color, border_color, icon=""):
    """Caja de nota/advertencia con fondo de color."""
    table = doc.add_table(rows=1, cols=1)
    table.style = 'Table Grid'
    cell = table.rows[0].cells[0]
    set_cell_bg(cell, bg_color)
    set_cell_borders(cell, border_color)
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    para = cell.paragraphs[0]
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after  = Pt(0)
    if icon:
        run_icon = para.add_run(icon + "  ")
        style_run(run_icon, size=10, bold=True, color=border_color)
    run = para.add_run(text)
    style_run(run, size=10, color=C_TEXT, italic=False)
    add_space(doc, 6)
    return table

def step_table(doc, steps, accent):
    """
    Tabla de pasos numerados: N | Accion | Detalle
    steps: list of (accion, detalle)
    """
    W = [0.8, 4.5, 10.9]
    table = doc.add_table(rows=len(steps), cols=3)
    table.style = 'Table Grid'
    for i, (accion, detalle) in enumerate(steps):
        row = table.rows[i]
        # Columna N
        c0 = row.cells[0]
        c0.width = Cm(W[0])
        set_cell_bg(c0, accent)
        set_cell_borders(c0, accent)
        set_cell_margins(c0, top=100, bottom=100, left=100, right=80)
        c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(str(i+1))
        style_run(r0, size=11, bold=True, color=C_WHITE)

        # Columna Accion
        c1 = row.cells[1]
        c1.width = Cm(W[1])
        set_cell_bg(c1, C_GRAY)
        set_cell_borders(c1)
        set_cell_margins(c1)
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(accion)
        style_run(r1, size=10, bold=True, color=C_DARK)

        # Columna Detalle
        c2 = row.cells[2]
        c2.width = Cm(W[2])
        set_cell_bg(c2, C_WHITE)
        set_cell_borders(c2)
        set_cell_margins(c2)
        p2 = c2.paragraphs[0]
        r2 = p2.add_run(detalle)
        style_run(r2, size=10, color=C_TEXT)

    add_space(doc, 6)

def role_cover_block(doc, rol, descripcion, color, light_color, icono_text):
    """Bloque de portada de sección para cada rol."""
    # Banda de color
    table = doc.add_table(rows=1, cols=1)
    table.style = 'Table Grid'
    cell = table.rows[0].cells[0]
    set_cell_bg(cell, color)
    set_cell_borders(cell, color)
    set_cell_margins(cell, top=280, bottom=280, left=400, right=200)

    p = cell.paragraphs[0]
    run_icono = p.add_run(icono_text + "  ")
    style_run(run_icono, size=32, bold=True, color=C_WHITE)

    run_rol = p.add_run(rol)
    style_run(run_rol, size=32, bold=True, color=C_WHITE)

    add_space(doc, 6)

    # Descripcion
    desc_tbl = doc.add_table(rows=1, cols=1)
    desc_tbl.style = 'Table Grid'
    dc = desc_tbl.rows[0].cells[0]
    set_cell_bg(dc, light_color)
    set_cell_borders(dc, color)
    set_cell_margins(dc, top=160, bottom=160, left=300, right=200)
    dp = dc.paragraphs[0]
    dr = dp.add_run(descripcion)
    style_run(dr, size=11, color=C_TEXT, italic=True)

    add_space(doc, 12)

def quick_ref_table(doc, rows_data, accent):
    """Tabla de referencia rapida: Accion | Ruta en el sistema"""
    W_cm = [5.5, 10.7]
    table = doc.add_table(rows=1 + len(rows_data), cols=2)
    table.style = 'Table Grid'

    # Header
    hr = table.rows[0]
    for j, hdr in enumerate(["Accion", "Ruta en el sistema"]):
        c = hr.cells[j]
        c.width = Cm(W_cm[j])
        set_cell_bg(c, accent)
        set_cell_borders(c, accent)
        set_cell_margins(c)
        p = c.paragraphs[0]
        r = p.add_run(hdr)
        style_run(r, size=9.5, bold=True, color=C_WHITE)

    for i, (accion, ruta) in enumerate(rows_data):
        row = table.rows[i+1]
        bg = C_GRAY if i % 2 == 0 else C_WHITE
        for j, val in enumerate([accion, ruta]):
            c = row.cells[j]
            c.width = Cm(W_cm[j])
            set_cell_bg(c, bg)
            set_cell_borders(c)
            set_cell_margins(c)
            p = c.paragraphs[0]
            r = p.add_run(val)
            style_run(r, size=10, color=C_TEXT, bold=(j==0))

    add_space(doc, 8)

# ══════════════════════════════════════════════════════════════════════════════
# DOCUMENTO
# ══════════════════════════════════════════════════════════════════════════════

doc = Document()

for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(2.5)
    section.right_margin  = Cm(2.5)

# ─── PORTADA GENERAL ─────────────────────────────────────────────────────────

add_space(doc, 60)

border_top = doc.add_paragraph()
add_paragraph_border_bottom(border_top, C_ACCENT, sz=16)
add_space(doc, 12)

add_para(doc, "RISTORANTINO", size=42, bold=True, color=C_DARK,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_before=20, space_after=4)
add_para(doc, "Sistema de Gestion de Restaurante", size=18, color=C_MID,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=20)

sub = add_para(doc, "Manual de Usuario — por Rol", size=15, italic=True, bold=True,
               color=C_ACCENT, align=WD_ALIGN_PARAGRAPH.CENTER,
               space_before=0, space_after=8)
add_paragraph_border_bottom(sub, "DDDDDD", sz=4)

add_space(doc, 40)

add_para(doc, "Version 1.0  |  Fecha: 29 de mayo de 2026", size=11, color="666666",
         align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4, space_after=2)
add_para(doc, "SG Systems  |  elvis.lopez@sgsystems.com.py", size=11, color="666666",
         align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=2)

add_space(doc, 50)
border_bot = doc.add_paragraph()
add_paragraph_border_bottom(border_bot, C_ACCENT, sz=16)

# Indice de roles
add_space(doc, 20)
add_para(doc, "Este manual contiene tres secciones independientes:", size=11, color=C_TEXT,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_before=10, space_after=6)

roles_list = [
    ("Manual 1", "MOZO",   "Apertura de mesas, toma de pedidos, gestion del salon"),
    ("Manual 2", "CHEF",   "Monitor de cocina, gestion de estados de preparacion"),
    ("Manual 3", "CAJERO", "Gestion de caja, cobro de pedidos, facturacion fiscal"),
]
for num, rol, desc in roles_list:
    t = doc.add_table(rows=1, cols=2)
    t.style = 'Table Grid'
    c0 = t.rows[0].cells[0]
    c1 = t.rows[0].cells[1]
    c0.width = Cm(3)
    c1.width = Cm(13.2)
    color_map = {"MOZO": C_GREEN, "CHEF": C_ORANGE, "CAJERO": C_PURPLE}
    clr = color_map[rol]
    set_cell_bg(c0, clr); set_cell_borders(c0, clr); set_cell_margins(c0)
    set_cell_bg(c1, C_GRAY); set_cell_borders(c1); set_cell_margins(c1)
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r0 = p0.add_run(rol)
    style_run(r0, size=11, bold=True, color=C_WHITE)
    p1 = c1.paragraphs[0]
    r1a = p1.add_run(num + " — ")
    style_run(r1a, size=10, bold=True, color=C_DARK)
    r1b = p1.add_run(desc)
    style_run(r1b, size=10, color=C_TEXT)
    add_space(doc, 3)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
# MANUAL 1 — MOZO
# ══════════════════════════════════════════════════════════════════════════════

add_heading(doc, "Manual 1 — Mozo", 1, color=C_GREEN, border_color=C_GREEN)
add_space(doc, 4)
role_cover_block(doc, "MOZO", (
    "El mozo es el nexo entre el cliente y la cocina. "
    "Gestiona el salon: abre mesas, toma pedidos, realiza modificaciones "
    "y coordina la entrega de platos. Este manual cubre todas las "
    "operaciones del dia a dia en el salon."
), C_GREEN, C_GREEN_L, "M")

# 1.1 Inicio de sesion
add_heading(doc, "1.1 Inicio de Sesion", 2, color=C_GREEN)
add_para(doc, "Para ingresar al sistema:")
step_table(doc, [
    ("Abrir navegador",    "Ingresar a http://localhost:8000 desde cualquier dispositivo de la red del restaurante."),
    ("Ingresar usuario",   "Escribir el nombre de usuario asignado por el administrador (ej: mozo01)."),
    ("Ingresar contrasena","Escribir la contrasena. Si la olvidaste, pedile al administrador que la reestablezca."),
    ("Clic en Iniciar",    "Al ingresar correctamente seras redirigido automaticamente al Salon Principal."),
], C_GREEN)

add_warning_box(doc,
    "IMPORTANTE: No compartas tu usuario y contrasena con otros empleados. "
    "Cada accion queda registrada con tu usuario para fines de auditoria.",
    "FFF3E0", C_ORANGE, "!")

# 1.2 Pantalla del salon
add_heading(doc, "1.2 Pantalla del Salon", 2, color=C_GREEN)
add_para(doc, (
    "Al ingresar veras el Salon Principal: una grilla con todas las mesas del restaurante "
    "organizadas por sector. Cada mesa muestra su numero y estado visual:"
))
add_space(doc)
color_states = doc.add_table(rows=5, cols=2)
color_states.style = 'Table Grid'
states = [
    ("LIBRE",     "Verde",    "La mesa esta disponible. Podes abrirla para un nuevo cliente.", "1A6B3C"),
    ("OCUPADA",   "Rojo",     "La mesa tiene un pedido activo en curso.", "B02020"),
    ("RESERVADA", "Amarillo", "La mesa tiene una reserva pero aun no hay pedido.", "A07800"),
    ("SUCIA",     "Gris",     "La mesa fue cerrada y necesita limpieza antes de usarse.", "555555"),
]
# Header
hr = color_states.rows[0]
for j, h in enumerate(["Estado", "Descripcion y accion posible"]):
    c = hr.cells[j]
    set_cell_bg(c, C_GREEN); set_cell_borders(c, C_GREEN); set_cell_margins(c)
    p = c.paragraphs[0]; r = p.add_run(h)
    style_run(r, size=9.5, bold=True, color=C_WHITE)

for i, (estado, color_nombre, desc, clr) in enumerate(states):
    row = color_states.rows[i+1]
    c0 = row.cells[0]; c1 = row.cells[1]
    set_cell_bg(c0, clr); set_cell_borders(c0); set_cell_margins(c0)
    set_cell_bg(c1, C_GRAY if i%2==0 else C_WHITE); set_cell_borders(c1); set_cell_margins(c1)
    p0 = c0.paragraphs[0]; r0 = p0.add_run(estado)
    style_run(r0, size=9.5, bold=True, color=C_WHITE)
    p1 = c1.paragraphs[0]; r1 = p1.add_run(desc)
    style_run(r1, size=10, color=C_TEXT)
add_space(doc, 8)

# 1.3 Abrir una mesa
add_heading(doc, "1.3 Abrir una Mesa", 2, color=C_GREEN)
add_para(doc, "Cuando llegan clientes y necesitas asignarles una mesa:")
step_table(doc, [
    ("Identificar mesa libre", "Busca en la grilla del salon una mesa de color VERDE (Libre) en el sector que prefieras."),
    ("Clic en la mesa",        "Hace clic sobre la tarjeta de la mesa. Se abrira un modal de confirmacion."),
    ("Confirmar apertura",     "Clic en 'Abrir Mesa'. El sistema crea automaticamente un pedido y cambia el estado de la mesa a OCUPADA (rojo)."),
    ("Pantalla de comanda",    "Seras redirigido directamente a la pantalla de toma de pedidos para esa mesa."),
], C_GREEN)

add_warning_box(doc,
    "Si la mesa ya esta OCUPADA y clicas en ella, el sistema te redirige directamente "
    "al pedido activo existente. Nunca se crean dos pedidos para la misma mesa.",
    C_GREEN_L, C_GREEN)

# 1.4 Tomar el pedido
add_heading(doc, "1.4 Tomar el Pedido (Comanda)", 2, color=C_GREEN)
add_para(doc, (
    "La pantalla de comanda esta dividida en dos columnas. "
    "A la izquierda el menu disponible, a la derecha el resumen del pedido actual."
))
add_space(doc, 4)
add_para(doc, "Como agregar productos al pedido:", bold=True, size=10.5)
step_table(doc, [
    ("Ver el menu",          "En la columna izquierda veras las categorias del menu (Entradas, Platos, Bebidas, etc.)."),
    ("Seleccionar categoria","Clic en una categoria para expandirla y ver sus productos con precio."),
    ("Agregar producto",     "Clic en el boton '+' junto al producto. El sistema lo agrega instantaneamente al pedido de la derecha."),
    ("Repetir",              "Para agregar mas unidades del mismo producto volvé a hacer clic en '+'. La cantidad se incrementa automaticamente."),
    ("Agregar nota",         "Si el cliente pide una modificacion (ej: 'sin cebolla'), clica el icono de nota en el item y escribi la observacion."),
], C_GREEN)

add_para(doc, "Como modificar cantidades o eliminar productos:", bold=True, size=10.5)
step_table(doc, [
    ("Aumentar cantidad",  "Clic en el boton '+' del item en la columna derecha del pedido."),
    ("Disminuir cantidad", "Clic en el boton '-' del item. Si llega a 0, el item se elimina del pedido."),
    ("Eliminar item",      "Clic en el boton de basura (papelera) junto al item en el pedido."),
], C_GREEN)

add_warning_box(doc,
    "Todos los cambios se guardan automaticamente. No hay boton 'Guardar'. "
    "El total se actualiza en tiempo real cada vez que agregues o quites un producto.",
    C_GREEN_L, C_GREEN)

# 1.5 Confirmar y volver al salon
add_heading(doc, "1.5 Volver al Salon", 2, color=C_GREEN)
add_para(doc, (
    "Cuando termines de cargar el pedido de una mesa y necesites atender otra:"
))
step_table(doc, [
    ("Boton 'Volver al Salon'", "Clic en el boton de volver. Si el pedido tiene items, la mesa queda activa y volveras al salon."),
    ("Pedido vacio",            "Si no agregaste ningun producto, el sistema cancela el pedido automaticamente y libera la mesa (queda Libre)."),
], C_GREEN)

# 1.6 Marcar mesa como sucia
add_heading(doc, "1.6 Marcar Mesa como Sucia", 2, color=C_GREEN)
add_para(doc, (
    "Cuando un cliente se va despues de pagar y la mesa necesita limpieza:"
))
step_table(doc, [
    ("Ir al salon",         "La mesa ya deberia estar en estado SUCIA (gris) tras el cobro del cajero."),
    ("Clic en la mesa",     "Podes acceder al modal de la mesa para ver su estado."),
    ("Marcar manualmente",  "Si necesitas marcarla como sucia vos mismo, usa el boton 'Marcar como Sucia' en el modal."),
    ("Notificar al personal","Avisa al personal de limpieza. Una vez limpia, el encargado la vuelve a Libre desde el admin."),
], C_GREEN)

# 1.7 Reservaciones
add_heading(doc, "1.7 Reservaciones", 2, color=C_GREEN)
add_para(doc, "Para gestionar reservas de clientes:")
step_table(doc, [
    ("Ver reservaciones",    "En el menu superior, ir a Reservaciones > Lista de Reservaciones."),
    ("Nueva reservacion",    "En el salon, clic en una mesa LIBRE y elegir 'Reservar'. Completa el formulario con datos del cliente, fecha y hora."),
    ("Activar reservacion",  "Cuando el cliente llega, clic en la mesa RESERVADA y luego en 'Activar Mesa'. Esto crea el pedido y la mesa pasa a Ocupada."),
    ("Cancelar reservacion", "En la lista de reservaciones, clic en 'Cancelar' junto a la reserva deseada."),
], C_GREEN)

# 1.8 Transferir mesa
add_heading(doc, "1.8 Transferir un Pedido a Otra Mesa", 2, color=C_GREEN)
add_para(doc, "Si los clientes necesitan cambiarse de mesa:")
step_table(doc, [
    ("Ir a la comanda",      "Clic en la mesa ocupada para abrir la comanda del pedido."),
    ("Boton Transferir",     "Busca el boton 'Transferir Mesa' en la pantalla de comanda."),
    ("Seleccionar destino",  "Elige la mesa destino (debe estar Libre). El pedido se mueve y la mesa origen queda Libre."),
], C_GREEN)

# 1.9 Referencia rapida
add_heading(doc, "1.9 Referencia Rapida — Mozo", 2, color=C_GREEN)
quick_ref_table(doc, [
    ("Ver el salon con las mesas",          "Salon Principal (pantalla de inicio)"),
    ("Abrir una mesa para nuevos clientes", "Salon > clic en mesa VERDE > Abrir Mesa"),
    ("Agregar productos al pedido",         "Comanda > columna izquierda > boton '+'"),
    ("Agregar nota a un item",              "Comanda > icono de nota junto al item"),
    ("Modificar o quitar un producto",      "Comanda > columna derecha > botones +/-/papelera"),
    ("Volver al salon sin cerrar el pedido","Comanda > boton Volver al Salon"),
    ("Crear una reservacion",               "Salon > clic en mesa LIBRE > Reservar"),
    ("Activar mesa con reserva",            "Salon > clic en mesa AMARILLA > Activar Mesa"),
    ("Transferir pedido a otra mesa",       "Comanda > boton Transferir Mesa"),
], C_GREEN)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
# MANUAL 2 — CHEF
# ══════════════════════════════════════════════════════════════════════════════

add_heading(doc, "Manual 2 — Chef", 1, color=C_ORANGE, border_color=C_ORANGE)
add_space(doc, 4)
role_cover_block(doc, "CHEF", (
    "El chef y el personal de cocina trabajan desde el Monitor de Cocina (KDS). "
    "Esta pantalla muestra en tiempo real todos los items de pedidos pendientes "
    "de preparacion, ordenados del mas antiguo al mas nuevo (FIFO). "
    "Este manual explica como usar el monitor y gestionar el flujo de cocina."
), C_ORANGE, C_ORANGE_L, "C")

# 2.1 Inicio de sesion
add_heading(doc, "2.1 Inicio de Sesion", 2, color=C_ORANGE)
add_para(doc, "El acceso al Monitor de Cocina requiere un usuario con rol Chef o Administrador.")
step_table(doc, [
    ("Abrir navegador",   "Ingresar a http://localhost:8000 desde la tablet o pantalla de la cocina."),
    ("Iniciar sesion",    "Usar el usuario y contrasena asignados por el administrador (ej: chef01)."),
    ("Ir a Cocina",       "En el menu superior clic en 'Monitor de Cocina', o ir directamente a /cocina/."),
], C_ORANGE)

add_warning_box(doc,
    "El Monitor de Cocina se actualiza automaticamente cada 20 segundos. "
    "No es necesario recargar la pagina manualmente.",
    C_ORANGE_L, C_ORANGE)

# 2.2 Monitor de cocina
add_heading(doc, "2.2 El Monitor de Cocina (KDS)", 2, color=C_ORANGE)
add_para(doc, (
    "El monitor muestra tarjetas de items. Cada tarjeta representa un plato o bebida "
    "de un pedido activo. Las tarjetas se ordenan del mas antiguo al mas nuevo "
    "para que los pedidos que entraron primero se atiendan primero (sistema FIFO)."
))
add_space(doc, 4)
add_para(doc, "Informacion que muestra cada tarjeta:", bold=True, size=10.5)
add_bullet(doc, "Numero de mesa: la mesa que hizo el pedido.")
add_bullet(doc, "Nombre del producto: que plato o bebida hay que preparar.")
add_bullet(doc, "Cantidad: cuantas unidades de ese producto hay que hacer.")
add_bullet(doc, "Observaciones: notas del mozo (ej: 'sin sal', 'bien cocido', 'alergia al gluten').")
add_bullet(doc, "Estado actual: Pendiente (naranja) o En Cocina (azul).")
add_bullet(doc, "Tiempo transcurrido: hace cuantos minutos entro el pedido.")
add_space(doc, 4)

add_para(doc, "Estados de los items en el monitor:", bold=True, size=10.5)
st_table = doc.add_table(rows=3, cols=3)
st_table.style = 'Table Grid'
st_headers = ["Estado", "Color", "Significado"]
hr2 = st_table.rows[0]
for j, h in enumerate(st_headers):
    c = hr2.cells[j]
    set_cell_bg(c, C_ORANGE); set_cell_borders(c, C_ORANGE); set_cell_margins(c)
    p = c.paragraphs[0]; r = p.add_run(h)
    style_run(r, size=9.5, bold=True, color=C_WHITE)
st_data = [
    ("Pendiente", "Naranja / Amarillo", "El item fue cargado por el mozo y esta esperando que cocina lo tome."),
    ("En Cocina", "Azul",               "El chef acepto el item y lo esta preparando actualmente."),
]
for i, (st, col, desc) in enumerate(st_data):
    row2 = st_table.rows[i+1]
    bg = C_GRAY if i%2==0 else C_WHITE
    for j, val in enumerate([st, col, desc]):
        c = row2.cells[j]
        set_cell_bg(c, bg); set_cell_borders(c); set_cell_margins(c)
        p = c.paragraphs[0]; r = p.add_run(val)
        style_run(r, size=10, color=C_TEXT)
add_space(doc, 8)

# 2.3 Flujo de trabajo
add_heading(doc, "2.3 Flujo de Trabajo en Cocina", 2, color=C_ORANGE)
add_para(doc, (
    "Cada item pasa por tres etapas bajo control del chef. "
    "El flujo completo es:"
))
add_space(doc, 4)

flow_p = add_para(doc, "PENDIENTE  -->  EN COCINA  -->  LISTO", size=13, bold=True, color=C_ORANGE,
                  align=WD_ALIGN_PARAGRAPH.CENTER)

add_space(doc, 8)
add_para(doc, "Paso a paso en la operativa diaria:", bold=True, size=10.5)
step_table(doc, [
    ("Revisar nuevos items",
     "Los items PENDIENTES aparecen en el monitor con una tarjeta de color naranja/amarillo. "
     "Leer nombre, cantidad y observaciones."),
    ("Empezar a preparar",
     "Clic en el boton 'Empezar' de la tarjeta. El item pasa a estado EN COCINA (azul). "
     "Esto le indica al mozo que el plato esta en preparacion."),
    ("Preparar el plato",
     "Preparar el plato segun las observaciones indicadas. El tiempo transcurrido sigue contando."),
    ("Marcar como Listo",
     "Cuando el plato esta listo para servir, clic en el boton 'Listo'. "
     "El item desaparece del monitor. El mozo recibe la notificacion para llevar el plato a la mesa."),
], C_ORANGE)

add_warning_box(doc,
    "ATENCION: Siempre leer las observaciones del item antes de empezar a prepararlo. "
    "Las notas de alergias y modificaciones del cliente estan ahi.",
    "FFF3E0", C_ORANGE, "!")

add_warning_box(doc,
    "El orden FIFO garantiza que los pedidos mas antiguos aparezcan primero. "
    "Trata de atender las tarjetas de arriba antes que las de abajo.",
    C_ORANGE_L, C_ORANGE)

# 2.4 Auto-refresco
add_heading(doc, "2.4 Actualizacion Automatica", 2, color=C_ORANGE)
add_para(doc, (
    "El monitor se recarga automaticamente cada 20 segundos sin necesidad de recargar la pagina. "
    "Veras aparecer nuevas tarjetas a medida que los mozos agregan items a los pedidos."
))
add_bullet(doc, "Nuevos items: aparecen automaticamente con estado PENDIENTE.")
add_bullet(doc, "Items completados: desaparecen del monitor cuando se marcan como LISTO.")
add_bullet(doc, "Hora del servidor: visible en pantalla para verificar sincronizacion.")

# 2.5 Referencia rapida
add_heading(doc, "2.5 Referencia Rapida — Chef", 2, color=C_ORANGE)
quick_ref_table(doc, [
    ("Abrir el monitor de cocina",           "/cocina/ o menu > Monitor de Cocina"),
    ("Ver items pendientes de preparacion",  "Monitor > tarjetas color naranja (Pendiente)"),
    ("Aceptar un item y empezar a cocinarlo","Monitor > tarjeta > boton 'Empezar'"),
    ("Ver items que ya estoy preparando",    "Monitor > tarjetas color azul (En Cocina)"),
    ("Marcar un plato como listo",           "Monitor > tarjeta > boton 'Listo'"),
    ("Ver observaciones del mozo",           "Monitor > tarjeta > campo Observaciones"),
    ("Ver hace cuanto entro el pedido",      "Monitor > tarjeta > campo de tiempo transcurrido"),
], C_ORANGE)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
# MANUAL 3 — CAJERO
# ══════════════════════════════════════════════════════════════════════════════

add_heading(doc, "Manual 3 — Cajero", 1, color=C_PURPLE, border_color=C_PURPLE)
add_space(doc, 4)
role_cover_block(doc, "CAJERO", (
    "El cajero gestiona la caja del restaurante: abre y cierra turnos, "
    "procesa el cobro de pedidos y emite la documentacion tributaria (facturas). "
    "Este manual cubre el ciclo completo del turno de caja."
), C_PURPLE, C_PURPLE_L, "Ca")

# 3.1 Inicio de sesion
add_heading(doc, "3.1 Inicio de Sesion", 2, color=C_PURPLE)
step_table(doc, [
    ("Abrir navegador",    "Ingresar a http://localhost:8000 desde la caja registradora."),
    ("Iniciar sesion",     "Usar el usuario y contrasena de cajero asignados por el administrador."),
    ("Ir a Gestion de Caja","En el menu superior clic en 'Caja' para acceder al panel principal."),
], C_PURPLE)

# 3.2 Abrir caja
add_heading(doc, "3.2 Abrir la Caja (Inicio de Turno)", 2, color=C_PURPLE)
add_para(doc, (
    "Al inicio de cada turno el cajero debe abrir la caja. "
    "Sin caja abierta no es posible cobrar pedidos ni emitir facturas."
))
step_table(doc, [
    ("Ir a Caja",          "Menu superior > Caja, o ir directamente a /caja/."),
    ("Seleccionar caja",   "En el formulario de apertura, elegir la caja fisica que vas a usar."),
    ("Ingresar monto",     "Escribir el monto de apertura: el efectivo con el que arranca la caja (fondo inicial)."),
    ("Confirmar apertura", "Clic en 'Abrir Caja'. El sistema crea un Arqueo en estado Abierto y registra la hora."),
    ("Verificar estado",   "La pantalla de caja mostrara el arqueo activo con los datos del turno."),
], C_PURPLE)

add_warning_box(doc,
    "Solo puede haber un arqueo abierto por usuario al mismo tiempo. "
    "Si ya tenes una caja abierta, el sistema te lo avisara y no creara una nueva.",
    C_PURPLE_L, C_PURPLE)

# 3.3 Cobrar un pedido
add_heading(doc, "3.3 Cobrar un Pedido y Emitir Factura", 2, color=C_PURPLE)
add_para(doc, (
    "Cuando una mesa solicita la cuenta, el cajero procesa el cobro desde el sistema. "
    "Este proceso genera automaticamente la factura fiscal."
))
step_table(doc, [
    ("Ir al salon o al historial",
     "Podes acceder al pedido desde el Salon Principal (clic en la mesa OCUPADA) "
     "o desde Pedidos > Historial."),
    ("Abrir el pedido",
     "Clic en la mesa para ver el detalle del pedido activo."),
    ("Ir a Cobrar",
     "Clic en el boton 'Cobrar Pedido'. Seras redirigido a la pantalla de cobro."),
    ("Verificar los items",
     "Revisa que los productos y cantidades sean correctos. Verificar el total."),
    ("Seleccionar cliente",
     "Busca el cliente por nombre o RUC/CI usando el buscador. Si es un cliente ocasional sin datos, deja el campo en blanco (se usara 'Sin Nombre / Ocasional')."),
    ("Aplicar descuento (opcional)",
     "Si corresponde un descuento, ingresar el monto en el campo Descuento. El total se recalcula automaticamente."),
    ("Seleccionar metodo de pago",
     "Elegir entre Efectivo, Tarjeta de credito, Tarjeta de debito u otro metodo disponible."),
    ("Confirmar cobro",
     "Clic en 'Cobrar'. El sistema genera automaticamente la factura con el numero correlativo "
     "del timbrado, calcula el IVA (10% y 5% segun producto) y actualiza el arqueo."),
    ("Imprimir factura",
     "Seras redirigido a la vista de factura imprimible. Usa Ctrl+P para imprimir "
     "o guardar como PDF para enviar al cliente."),
], C_PURPLE)

add_warning_box(doc,
    "IMPORTANTE: El cobro SOLO funciona si tenes la caja abierta (Arqueo activo). "
    "Si intentes cobrar sin caja abierta, el sistema te redirigira a la pantalla de caja.",
    "FFF3E0", C_ORANGE, "!")

add_warning_box(doc,
    "El timbrado debe estar vigente (no vencido) y dentro del rango de numeracion autorizado. "
    "Si el sistema muestra un error de timbrado, notifica al administrador para que lo actualice.",
    C_PURPLE_L, C_PURPLE)

# 3.4 IVA y documentos tributarios
add_heading(doc, "3.4 Facturacion e IVA", 2, color=C_PURPLE)
add_para(doc, "El sistema calcula y desglosa el IVA automaticamente segun la tasa de cada producto:")
add_space(doc)
iva_table = doc.add_table(rows=4, cols=3)
iva_table.style = 'Table Grid'
iva_headers = ["Tasa IVA", "Calculo", "Ejemplo de productos"]
ihr = iva_table.rows[0]
for j, h in enumerate(iva_headers):
    c = ihr.cells[j]
    set_cell_bg(c, C_PURPLE); set_cell_borders(c, C_PURPLE); set_cell_margins(c)
    p = c.paragraphs[0]; r = p.add_run(h)
    style_run(r, size=9.5, bold=True, color=C_WHITE)

iva_data = [
    ("10%",    "IVA = total_linea / 11",  "Bebidas alcoholicas, servicios, preparaciones gastronomicas en general."),
    ("5%",     "IVA = total_linea / 21",  "Productos de la canasta familiar, ciertos alimentos basicos."),
    ("Exento", "Sin IVA aplicado",         "Productos exentos segun legislacion fiscal vigente."),
]
for i, (tasa, calc, ej) in enumerate(iva_data):
    row = iva_table.rows[i+1]
    bg = C_GRAY if i%2==0 else C_WHITE
    for j, val in enumerate([tasa, calc, ej]):
        c = row.cells[j]
        set_cell_bg(c, bg); set_cell_borders(c); set_cell_margins(c)
        p = c.paragraphs[0]; r = p.add_run(val)
        style_run(r, size=10, color=C_TEXT)
add_space(doc, 8)

add_para(doc, (
    "La factura generada incluye: numero de timbrado, numero correlativo, fecha y hora, "
    "datos del cliente, detalle de items, subtotales por tasa de IVA, descuento aplicado "
    "y total general."
))

# 3.5 Cerrar caja
add_heading(doc, "3.5 Cerrar la Caja (Cierre de Turno)", 2, color=C_PURPLE)
add_para(doc, (
    "Al finalizar el turno, el cajero debe cerrar la caja para registrar "
    "el arqueo de cierre y dejar el sistema listo para el proximo turno."
))
step_table(doc, [
    ("Ir a Gestion de Caja", "Menu superior > Caja."),
    ("Revisar estadisticas", "Verificar el total vendido y la cantidad de facturas del turno mostradas en el panel."),
    ("Contar el efectivo",   "Contar fisicamente el dinero en caja: monto de apertura + ventas en efectivo."),
    ("Ingresar monto real",  "En el formulario de cierre, ingresar el monto fisico contado."),
    ("Confirmar cierre",     "Clic en 'Cerrar Caja'. El sistema calcula la diferencia de arqueo "
                             "(diferencia = monto contado - monto esperado por sistema)."),
    ("Revisar diferencia",   "Si la diferencia es 0: todo cuadra. Si es negativa: hay faltante. "
                             "Si es positiva: hay sobrante. El resultado queda registrado."),
], C_PURPLE)

add_warning_box(doc,
    "Una vez cerrada la caja NO se puede reabrir ni modificar el arqueo. "
    "Asegurate de haber registrado todos los cobros antes de cerrar.",
    "FFF3E0", C_ORANGE, "!")

# 3.6 Reportes
add_heading(doc, "3.6 Reportes Disponibles", 2, color=C_PURPLE)
add_para(doc, "El cajero tiene acceso a los siguientes reportes desde el menu:")
add_space(doc)
rep_table = doc.add_table(rows=5, cols=2)
rep_table.style = 'Table Grid'
rep_headers = ["Reporte", "Descripcion"]
rhr = rep_table.rows[0]
for j, h in enumerate(rep_headers):
    c = rhr.cells[j]
    set_cell_bg(c, C_PURPLE); set_cell_borders(c, C_PURPLE); set_cell_margins(c)
    p = c.paragraphs[0]; r = p.add_run(h)
    style_run(r, size=9.5, bold=True, color=C_WHITE)
rep_data = [
    ("Historial de Pedidos", "Lista de todos los pedidos con filtros por fecha, estado y mesa. Permite ver el detalle de cada uno."),
    ("Reporte de IVA",       "Resumen de IVA desglosado por tasa (10%, 5%, Exento) para el periodo seleccionado."),
    ("Reporte Z",            "Cierre fiscal del dia con el total de ventas, cantidad de comprobantes y montos por metodo de pago."),
    ("Exportar a Excel",     "Tanto el historial como el dashboard pueden exportarse en formato .xlsx para analisis externo."),
]
for i, (rep, desc) in enumerate(rep_data):
    row = rep_table.rows[i+1]
    bg = C_GRAY if i%2==0 else C_WHITE
    c0 = row.cells[0]; c1 = row.cells[1]
    set_cell_bg(c0, bg); set_cell_borders(c0); set_cell_margins(c0)
    set_cell_bg(c1, bg); set_cell_borders(c1); set_cell_margins(c1)
    p0 = c0.paragraphs[0]; r0 = p0.add_run(rep); style_run(r0, size=10, bold=True, color=C_DARK)
    p1 = c1.paragraphs[0]; r1 = p1.add_run(desc); style_run(r1, size=10, color=C_TEXT)
add_space(doc, 8)

# 3.7 Gestion de clientes
add_heading(doc, "3.7 Gestion de Clientes", 2, color=C_PURPLE)
add_para(doc, "El cajero puede registrar y editar datos de clientes para la facturacion:")
step_table(doc, [
    ("Ir a Clientes",      "Menu superior > Clientes > Lista de Clientes."),
    ("Nuevo cliente",      "Clic en 'Nuevo Cliente'. Completar nombre/razon social, RUC/CI, telefono y email."),
    ("Editar cliente",     "En la lista, clic en el nombre del cliente y luego en 'Editar'."),
    ("Buscar al cobrar",   "En la pantalla de cobro, usar el campo de busqueda para encontrar al cliente por nombre o RUC/CI."),
], C_PURPLE)

# 3.8 Referencia rapida
add_heading(doc, "3.8 Referencia Rapida — Cajero", 2, color=C_PURPLE)
quick_ref_table(doc, [
    ("Abrir la caja al inicio del turno",    "Menu > Caja > Seleccionar caja > Ingresar monto > Abrir Caja"),
    ("Ver estadisticas del turno actual",    "Menu > Caja (panel principal con totales del arqueo)"),
    ("Cobrar un pedido y emitir factura",    "Salon > clic en mesa OCUPADA > Cobrar Pedido > Completar formulario > Cobrar"),
    ("Imprimir o guardar una factura",       "Post-cobro seras redirigido a la factura > Ctrl+P para imprimir"),
    ("Cerrar la caja al final del turno",    "Menu > Caja > Ingresar monto contado > Cerrar Caja"),
    ("Ver historial de pedidos",             "Menu > Pedidos > Historial de Pedidos"),
    ("Ver reporte de IVA del periodo",       "Menu > Reportes > Reporte de IVA"),
    ("Generar Reporte Z del dia",            "Menu > Reportes > Reporte Z"),
    ("Exportar datos a Excel",               "Historial o Dashboard > boton Exportar Excel"),
    ("Registrar o editar un cliente",        "Menu > Clientes > Lista de Clientes"),
], C_PURPLE)

# ── GUARDAR ───────────────────────────────────────────────────────────────────
output = r"C:\Proyecto\docs\manual_usuario_ristorantino.docx"
doc.save(output)
print(f"Documento generado: {output}")
