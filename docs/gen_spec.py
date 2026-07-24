"""
Genera la Especificacion Tecnica de Ristorantino como .docx
usando python-docx (no requiere Node.js).
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

# ── Paleta ───────────────────────────────────────────────────────────────────
C_DARK   = "1B2B3A"
C_MID    = "2E6DA4"
C_ACCENT = "C8A96E"
C_LIGHT  = "D5E8F0"
C_GRAY   = "F5F5F0"
C_TEXT   = "2C2C2C"
C_WHITE  = "FFFFFF"
C_CODE   = "F0F0EA"

def hex_to_rgb(h):
    return RGBColor(int(h[0:2],16), int(h[2:4],16), int(h[4:6],16))

# ── Helpers de formato ────────────────────────────────────────────────────────

def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def set_cell_borders(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), 'CCCCCC')
        tcBorders.append(el)
    tcPr.append(tcBorders)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for side, val in (('top', top), ('bottom', bottom), ('left', left), ('right', right)):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:w'), str(val))
        el.set(qn('w:type'), 'dxa')
        tcMar.append(el)
    tcPr.append(tcMar)

def add_paragraph_border_bottom(para, hex_color, sz=8):
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(sz))
    bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), hex_color)
    pBdr.append(bottom)
    pPr.append(pBdr)

def style_run(run, size=11, bold=False, italic=False, color=C_TEXT, font='Arial'):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = hex_to_rgb(color)

def add_heading(doc, text, level, color=None, border=False):
    para = doc.add_heading(text, level=level)
    para.clear()  # clear auto-added run
    if level == 1:
        sz, bold, c = 18, True, color or C_DARK
    elif level == 2:
        sz, bold, c = 14, True, color or C_MID
    else:
        sz, bold, c = 12, True, color or C_DARK
    run = para.add_run(text)
    style_run(run, size=sz, bold=bold, color=c)
    if border:
        add_paragraph_border_bottom(para, C_ACCENT, sz=8)
    return para

def add_para(doc, text, size=10.5, color=C_TEXT, bold=False, italic=False,
             align=WD_ALIGN_PARAGRAPH.LEFT, space_before=3, space_after=4):
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(space_before)
    para.paragraph_format.space_after = Pt(space_after)
    para.alignment = align
    run = para.add_run(text)
    style_run(run, size=size, bold=bold, italic=italic, color=color)
    return para

def add_bullet(doc, text):
    para = doc.add_paragraph(style='List Bullet')
    para.paragraph_format.space_before = Pt(2)
    para.paragraph_format.space_after = Pt(2)
    run = para.add_run(text)
    style_run(run, size=10, color=C_TEXT)
    return para

# ── Tablas ────────────────────────────────────────────────────────────────────

def make_table(doc, headers, rows, col_widths_cm, code_col=None):
    """
    headers: list of str
    rows: list of list of str
    col_widths_cm: list of float
    code_col: index of column to render in monospace (0-based), or None
    """
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'

    # Widths
    for i, row in enumerate(table.rows):
        for j, cell in enumerate(row.cells):
            cell.width = Cm(col_widths_cm[j])

    # Header row
    hdr_row = table.rows[0]
    for j, (cell, hdr) in enumerate(zip(hdr_row.cells, headers)):
        set_cell_bg(cell, C_DARK)
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        para = cell.paragraphs[0]
        run = para.add_run(hdr)
        style_run(run, size=9.5, bold=True, color=C_WHITE)

    # Data rows
    for i, data_row in enumerate(rows):
        row_obj = table.rows[i + 1]
        bg = C_GRAY if i % 2 == 0 else C_WHITE
        for j, (cell, val) in enumerate(zip(row_obj.cells, data_row)):
            cell_bg = C_CODE if code_col is not None and j == code_col else bg
            set_cell_bg(cell, cell_bg)
            set_cell_borders(cell)
            set_cell_margins(cell)
            para = cell.paragraphs[0]
            run = para.add_run(val)
            font = 'Courier New' if code_col is not None and j == code_col else 'Arial'
            style_run(run, size=9, color=C_TEXT, font=font)

    return table

def add_space(doc, size=6):
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(0)
    run = para.add_run()
    run.font.size = Pt(size)

# ── Documento ────────────────────────────────────────────────────────────────

doc = Document()

# Margenes de pagina
for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(2.5)
    section.right_margin  = Cm(2.5)

# ── PORTADA ───────────────────────────────────────────────────────────────────

add_space(doc, 60)

# Linea decorativa superior
l = doc.add_paragraph()
add_paragraph_border_bottom(l, C_ACCENT, sz=16)
add_space(doc, 12)

# Titulo
t = add_para(doc, "RISTORANTINO", size=40, bold=True, color=C_DARK,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_before=16, space_after=4)

add_para(doc, "Sistema de Gestion de Restaurante", size=18, color=C_MID,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=16)

sub = add_para(doc, "Especificacion Tecnica del Proyecto", size=14, italic=True,
               bold=True, color=C_ACCENT, align=WD_ALIGN_PARAGRAPH.CENTER,
               space_before=0, space_after=4)
add_paragraph_border_bottom(sub, "DDDDDD", sz=4)

add_space(doc, 40)

add_para(doc, "Version 1.0", size=11, color="666666",
         align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4, space_after=2)
add_para(doc, "Fecha: 29 de mayo de 2026", size=11, color="666666",
         align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=2)
add_para(doc, "Autor: SG Systems", size=11, color="666666",
         align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=2)
add_para(doc, "Contacto: elvis.lopez@sgsystems.com.py", size=11, color="666666",
         align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=2)

add_space(doc, 60)
l2 = doc.add_paragraph()
add_paragraph_border_bottom(l2, C_ACCENT, sz=16)

doc.add_page_break()

# ── 1. INTRODUCCION ───────────────────────────────────────────────────────────

add_heading(doc, "1. Introduccion", 1, border=True)

add_heading(doc, "1.1 Descripcion General", 2)
add_para(doc, (
    "Ristorantino es un sistema de gestion integral para restaurantes desarrollado con Django. "
    "Cubre las operaciones diarias del salon: administracion de mesas, toma de pedidos en comanda "
    "digital, display de cocina en tiempo real y emision de documentos tributarios con trazabilidad "
    "fiscal completa."
))
add_para(doc, (
    "El sistema esta disenado para personal de restaurante (mozos, cajeros, chef) con interfaces "
    "diferenciadas por rol, garantizando que cada empleado acceda unicamente a las funciones que le corresponden."
))

add_heading(doc, "1.2 Objetivos del Sistema", 2)
add_bullet(doc, "Digitalizar la operativa del salon: apertura de mesas, toma de pedidos y cierre de cuenta.")
add_bullet(doc, "Proporcionar un monitor de cocina con cola FIFO para optimizar tiempos de preparacion.")
add_bullet(doc, "Emitir documentos tributarios (facturas, recibos) con numero de timbrado y trazabilidad de caja.")
add_bullet(doc, "Registrar KPIs de margen por producto mediante seguimiento de precio de venta y costo de insumo.")
add_bullet(doc, "Mantener un historial de arqueos de caja para conciliacion contable.")

add_heading(doc, "1.3 Alcance", 2)
add_para(doc, "El sistema abarca los siguientes modulos funcionales:")
add_bullet(doc, "Gestion de mesas y sectores del restaurante.")
add_bullet(doc, "Ingreso y modificacion de pedidos (comanda digital).")
add_bullet(doc, "Flujo de cocina con cambio de estados por item.")
add_bullet(doc, "Facturacion fiscal con timbrados y tipos de documento.")
add_bullet(doc, "Administracion de caja y arqueos.")
add_bullet(doc, "Panel de administracion Django para configuracion del sistema.")

doc.add_page_break()

# ── 2. STACK TECNOLOGICO ──────────────────────────────────────────────────────

add_heading(doc, "2. Stack Tecnologico", 1, border=True)
add_para(doc, "La siguiente tabla resume las tecnologias utilizadas en el proyecto:")
add_space(doc)

make_table(doc,
    headers=["Componente", "Tecnologia", "Detalle"],
    rows=[
        ["Backend framework", "Django 6.0.3", "ORM, autenticacion, admin, rutas, plantillas"],
        ["Base de datos", "PostgreSQL 15+", "127.0.0.1:5432 / db: ristorantino / user: postgres"],
        ["Frontend CSS", "Bootstrap 5.3.0", "Grid responsivo, componentes y utilidades"],
        ["Tema visual", "CSS Japandi personalizado", "Variables CSS con paleta minimalista Japandi"],
        ["Servidor de desarrollo", "Django runserver", "http://localhost:8000"],
        ["Lenguaje", "Python 3.12+", "Entorno virtual en env\\ (Windows)"],
        ["Control de versiones", "Git", "Repositorio local"],
    ],
    col_widths_cm=[4.2, 4.5, 7.5]
)

add_space(doc)
add_heading(doc, "2.1 Requisitos del Entorno", 2)
add_bullet(doc, "Python 3.12 o superior instalado en el sistema.")
add_bullet(doc, "PostgreSQL 15+ con una base de datos llamada 'ristorantino' creada previamente.")
add_bullet(doc, "Node.js no es requerido (Bootstrap se sirve via CDN).")
add_bullet(doc, "Sistema operativo Windows (el entorno virtual usa Scripts\\ en lugar de bin\\).")

doc.add_page_break()

# ── 3. ARQUITECTURA ───────────────────────────────────────────────────────────

add_heading(doc, "3. Arquitectura del Sistema", 1, border=True)

add_heading(doc, "3.1 Patron de Diseno", 2)
add_para(doc, (
    "El proyecto sigue el patron MVT (Model-View-Template) de Django, variante del MVC clasico:"
))
add_bullet(doc, "Model: Define la estructura de datos y logica de negocio (modelos ORM en pedidos/models.py).")
add_bullet(doc, "View: Procesa peticiones HTTP y orquesta la logica (vistas en pedidos/views.py).")
add_bullet(doc, "Template: Genera el HTML final combinando datos con plantillas (pedidos/templates/).")

add_heading(doc, "3.2 Estructura de Directorios", 2)
add_para(doc, "La siguiente estructura describe la organizacion del proyecto:")
add_space(doc)

make_table(doc,
    headers=["Ruta", "Descripcion"],
    rows=[
        ["ristorantino/",              "Raiz del proyecto"],
        ["  core/",                    "Configuracion global del proyecto Django"],
        ["    settings.py",            "Parametros de BD, apps instaladas, templates, static"],
        ["    urls.py",                "Enrutamiento principal (incluye pedidos.urls)"],
        ["    wsgi.py / asgi.py",      "Entry points de servidor"],
        ["  pedidos/",                 "Aplicacion principal con toda la logica de negocio"],
        ["    models.py",              "Definicion de todos los modelos ORM"],
        ["    views.py",               "Vistas HTTP y endpoints AJAX"],
        ["    urls.py",                "Rutas de la aplicacion pedidos"],
        ["    admin.py",               "Registro de modelos en el panel de administracion"],
        ["    templates/",             "Plantillas HTML de la aplicacion"],
        ["      base.html",            "Layout base (Bootstrap 5.3 + tema Japandi)"],
        ["      salon.html",           "Grilla de mesas con visualizacion de sillas"],
        ["      comanda.html",         "Pantalla de toma de pedidos (dos columnas)"],
        ["  static/",                  "Archivos estaticos del proyecto"],
        ["    css/",                   "Hojas de estilo personalizadas (tema Japandi)"],
        ["    js/",                    "Scripts JavaScript del frontend"],
        ["    img/",                   "Imagenes y recursos graficos"],
        ["  env/",                     "Entorno virtual Python (no versionado)"],
        ["  requirements.txt",         "Dependencias Python del proyecto"],
        ["  manage.py",                "Utilidad de administracion de Django"],
    ],
    col_widths_cm=[5.5, 10.7],
    code_col=0
)

add_space(doc)
add_heading(doc, "3.3 Configuracion de Base de Datos", 2)
add_para(doc, "Las credenciales de conexion a PostgreSQL estan configuradas en core/settings.py:")
add_bullet(doc, "Motor: django.db.backends.postgresql")
add_bullet(doc, "Host: 127.0.0.1  |  Puerto: 5432")
add_bullet(doc, "Nombre de base de datos: ristorantino")
add_bullet(doc, "Usuario: postgres  /  Contrasena: admin")
add_para(doc, (
    "Tambien existe un archivo db.sqlite3 en el directorio raiz para uso local sin PostgreSQL, "
    "aunque la configuracion de settings.py apunta a Postgres por defecto."
))

doc.add_page_break()

# ── 4. MODELO DE DATOS ────────────────────────────────────────────────────────

add_heading(doc, "4. Modelo de Datos", 1, border=True)
add_para(doc, (
    "El modelo de datos esta organizado en cinco grupos funcionales que reflejan los distintos "
    "dominios del negocio restaurantero."
))

add_heading(doc, "4.1 Acceso y Personal", 2)
make_table(doc,
    headers=["Modelo", "Campos clave", "Descripcion"],
    rows=[
        ["Rol",      "nombre (str)",
         "Define el rol del empleado. Valores tipicos: mozo, cajero, chef."],
        ["Empleado", "user (OneToOne -> User), rol (FK -> Rol)",
         "Extiende el User de Django con un rol. Acceso: user.empleado.rol.nombre"],
    ],
    col_widths_cm=[3.2, 5.0, 8.0]
)

add_space(doc)
add_heading(doc, "4.2 Infraestructura y Menu", 2)
make_table(doc,
    headers=["Modelo", "Campos clave", "Descripcion"],
    rows=[
        ["Sector",   "nombre (str)",
         "Agrupa mesas por zona del restaurante (Ej: salon principal, terraza, privado)."],
        ["Mesa",     "numero, sector (FK), estado (choices)",
         "Unidad fisica del restaurante. Cuatro estados: Libre, Ocupada, Reservada, Sucia."],
        ["Categoria","nombre (str)",
         "Clasificacion de productos del menu (Ej: Entradas, Platos principales, Bebidas)."],
        ["Producto", "nombre, categoria (FK), precio_venta, costo_insumo, activo",
         "Item del menu. Almacena precio_venta y costo_insumo para calcular KPIs de margen bruto."],
    ],
    col_widths_cm=[3.2, 5.0, 8.0]
)

add_space(doc)
add_heading(doc, "4.3 Pedidos", 2)
make_table(doc,
    headers=["Modelo", "Campos clave", "Descripcion"],
    rows=[
        ["Cliente",       "nombre, telefono, email",
         "Informacion del cliente. Un cliente por defecto se asigna al abrir una mesa sin datos."],
        ["Pedido",        "mesa (FK), cliente (FK), estado, total, fecha_apertura",
         "Pedido activo en una mesa. actualizar_total() usa F-expressions para sumar cantidad * precio_historico."],
        ["PedidoDetalle", "pedido (FK), producto (FK), cantidad, precio_historico",
         "Linea de detalle. precio_historico es un snapshot del precio al momento del pedido."],
    ],
    col_widths_cm=[3.2, 5.0, 8.0]
)

add_space(doc)
add_heading(doc, "4.4 Cocina", 2)
make_table(doc,
    headers=["Modelo", "Campos clave", "Descripcion"],
    rows=[
        ["FlujoEstado",
         "detalle (FK -> PedidoDetalle), estado (choices), timestamp_creacion, timestamp_cambio",
         "Rastrea cada item de pedido a traves de los estados de cocina. Permite cola FIFO ordenando por timestamp_creacion ascendente."],
    ],
    col_widths_cm=[3.2, 5.0, 8.0]
)

add_space(doc)
add_heading(doc, "4.5 Contabilidad y Documentos Tributarios", 2)
make_table(doc,
    headers=["Modelo", "Campos clave", "Descripcion"],
    rows=[
        ["TipoDocumento",     "nombre, codigo (str)",
         "Define los tipos de comprobante fiscal (Factura, Recibo, Nota de credito, etc.)."],
        ["Timbrado",          "numero, fecha_inicio, fecha_fin, activo",
         "Numero de timbrado habilitado por la autoridad fiscal. Controla la vigencia de emision."],
        ["Caja",              "nombre, saldo_inicial, activo",
         "Representa una caja registradora fisica. Puede asociarse a multiples timbrados."],
        ["CajaTimbrado",      "caja (FK), timbrado (FK), numero_desde, numero_hasta, ultimo_numero",
         "Vincula una caja a un timbrado con el rango de numeracion autorizado."],
        ["Arqueo",            "caja (FK), fecha, saldo_apertura, saldo_cierre, responsable (FK)",
         "Registro de conciliacion de caja al inicio y cierre del turno."],
        ["DocumentoTributario","tipo (FK), caja_timbrado (FK), numero, fecha, total, cliente (FK)",
         "Comprobante fiscal emitido. Vincula el pedido al timbrado, numero correlativo y cliente."],
        ["PedidoDocumento",   "pedido (FK), documento (FK)",
         "Tabla puente que vincula un Pedido con su DocumentoTributario correspondiente."],
    ],
    col_widths_cm=[3.8, 5.4, 7.0]
)

doc.add_page_break()

# ── 5. ESTADOS DE MESA ────────────────────────────────────────────────────────

add_heading(doc, "5. Estados de Mesa", 1, border=True)
add_para(doc, "Cada mesa transita por los siguientes estados durante su ciclo de vida operativo:")
add_space(doc)

make_table(doc,
    headers=["Estado", "Significado", "Comportamiento del sistema"],
    rows=[
        ["Libre",     "Mesa disponible para ser asignada",
         "Estado inicial. El boton 'Abrir Mesa' esta habilitado."],
        ["Ocupada",   "Mesa con un pedido activo en curso",
         "Generado por abrir_mesa(). Redirige si ya existe pedido activo."],
        ["Reservada", "Mesa con reserva pero sin pedido activo",
         "Gestionado manualmente o por modulo de reservas futuro."],
        ["Sucia",     "Mesa liberada pendiente de limpieza",
         "Transicion post-cierre de pedido. Requiere accion del personal para liberar."],
    ],
    col_widths_cm=[3.2, 5.0, 8.0]
)

add_space(doc)
add_para(doc, "Diagrama de transicion de estados:", bold=True, size=10.5)
add_space(doc, 3)
t = add_para(doc,
    "Libre  ->  [abrir_mesa()]  ->  Ocupada  ->  [cerrar pedido]  ->  Sucia  ->  [limpiar]  ->  Libre",
    size=9.5, color=C_DARK)
t.runs[0].font.name = 'Courier New'
add_para(doc, (
    "La transicion a Reservada puede originarse desde Libre, gestionada por el modulo de reservas "
    "o por accion manual del administrador."
))

doc.add_page_break()

# ── 6. FLUJO DE COCINA ────────────────────────────────────────────────────────

add_heading(doc, "6. Flujo de Cocina", 1, border=True)

add_heading(doc, "6.1 Estados del FlujoEstado", 2)
add_para(doc, (
    "Cada item (PedidoDetalle) genera un registro FlujoEstado que avanza secuencialmente. "
    "La vista monitor_cocina() muestra los items ordenados por timestamp_creacion ascendente (FIFO) "
    "para que los pedidos mas antiguos se atiendan primero."
))
add_space(doc)

make_table(doc,
    headers=["Estado origen", "Estado destino", "Evento", "Notas"],
    rows=[
        ["Pendiente", "Cocina",    "Empleado de cocina acepta el item",
         "FlujoEstado creado al agregar item al pedido."],
        ["Cocina",    "Listo",     "El plato esta preparado",
         "Se registra timestamp_cambio. Visible en monitor."],
        ["Listo",     "Entregado", "Mozo entrega el plato en mesa",
         "Item sale del monitor de cocina."],
        ["Entregado", "(final)",   "Estado terminal del item",
         "Se conserva para auditoria y KPIs."],
    ],
    col_widths_cm=[3.2, 3.2, 4.0, 5.8]
)

add_space(doc)
add_heading(doc, "6.2 Diagrama de Flujo", 2)
lines = [
    "[Item creado]",
    "     |",
    "     v",
    " PENDIENTE  -->  COCINA  -->  LISTO  -->  ENTREGADO",
    "              (chef)       (chef)      (mozo)",
]
for line in lines:
    p2 = add_para(doc, line, size=9.5, color=C_DARK, space_before=0, space_after=0)
    p2.runs[0].font.name = 'Courier New'

add_space(doc)
add_heading(doc, "6.3 Logica de cambiar_estado_cocina()", 2)
add_bullet(doc, "Recibe el ID del FlujoEstado y determina el siguiente estado en la secuencia.")
add_bullet(doc, "Actualiza el campo estado y registra el timestamp_cambio con la hora actual.")
add_bullet(doc, "Retorna una respuesta JSON con el nuevo estado para actualizar la UI sin recargar.")
add_bullet(doc, "Los items en estado 'Entregado' desaparecen del monitor de cocina.")

doc.add_page_break()

# ── 7. VISTAS Y ENDPOINTS ─────────────────────────────────────────────────────

add_heading(doc, "7. Vistas y Endpoints", 1, border=True)
add_para(doc, (
    "Las siguientes vistas componen la interfaz funcional del sistema. "
    "Todas residen en pedidos/views.py y estan enrutadas desde pedidos/urls.py."
))
add_space(doc)

make_table(doc,
    headers=["Vista / Funcion", "URL / Metodo", "Descripcion"],
    rows=[
        ["salon_principal()", "GET /salon/",
         "Grilla de mesas por sector con modales para acciones rapidas."],
        ["abrir_mesa()", "POST /mesa/abrir/",
         "Abre una mesa, crea Pedido con cliente por defecto; redirige si ya esta ocupada."],
        ["tomar_pedido()", "GET/POST /pedido/<id>/",
         "Pantalla de ingreso de pedidos (comanda.html), layout dos columnas menu vs. pedido."],
        ["agregar_producto_ajax()", "POST /pedido/agregar/",
         "Endpoint AJAX; incrementa cantidad si el item existe o crea nuevo PedidoDetalle; retorna JSON."],
        ["monitor_cocina()", "GET /cocina/",
         "Display de cocina; items ordenados por mas antiguos primero (FIFO)."],
        ["cambiar_estado_cocina()", "POST /cocina/estado/",
         "Avanza FlujoEstado para un item de pedido."],
    ],
    col_widths_cm=[4.5, 3.5, 8.2],
    code_col=0
)

add_space(doc)
add_heading(doc, "7.1 Endpoint AJAX: agregar_producto_ajax()", 2)
add_para(doc, "Este endpoint es el nucleo de la comanda digital. Su logica de negocio es:")
add_bullet(doc, "Recibe el ID del pedido y el ID del producto via POST.")
add_bullet(doc, "Busca si ya existe un PedidoDetalle para ese producto en el pedido activo.")
add_bullet(doc, "Si existe: incrementa el campo 'cantidad' en 1 y llama a actualizar_total().")
add_bullet(doc, "Si no existe: crea un nuevo PedidoDetalle y crea el FlujoEstado inicial en 'Pendiente'.")
add_bullet(doc, "Retorna un objeto JSON con los datos actualizados para refrescar la UI sin recargar.")

doc.add_page_break()

# ── 8. FRONTEND ───────────────────────────────────────────────────────────────

add_heading(doc, "8. Frontend y Plantillas", 1, border=True)

add_heading(doc, "8.1 Layout Base (base.html)", 2)
add_para(doc, "Todas las vistas heredan de base.html, que incluye:")
add_bullet(doc, "Bootstrap 5.3.0 via CDN para el sistema de grilla, componentes y utilidades.")
add_bullet(doc, "Variables CSS personalizadas del tema Japandi (paleta minimalista con tonos neutros).")
add_bullet(doc, "Barra de navegacion con acceso rapido segun el rol del usuario autenticado.")
add_bullet(doc, "Bloque de mensajes Django para notificaciones de acciones (alertas Bootstrap).")

add_heading(doc, "8.2 Salon Principal (salon.html)", 2)
add_para(doc, "La vista del salon utiliza CSS Grid para renderizar las mesas con visualizacion de sillas:")
add_bullet(doc, "Cada mesa se representa como una tarjeta con color segun su estado (verde=Libre, rojo=Ocupada, amarillo=Reservada, gris=Sucia).")
add_bullet(doc, "Las sillas se dibujan alrededor de la mesa usando posicionamiento CSS absoluto.")
add_bullet(doc, "Modales Bootstrap permiten acciones rapidas (abrir mesa, ver pedido) sin navegar a otra pagina.")

add_heading(doc, "8.3 Comanda (comanda.html)", 2)
add_para(doc, "La pantalla de toma de pedidos usa un layout de dos columnas:")
add_bullet(doc, "Columna izquierda: catalogo de productos agrupados por categoria, con boton de agregar.")
add_bullet(doc, "Columna derecha: resumen del pedido actual con cantidades, precios historicos y total.")
add_bullet(doc, "La interaccion usa fetch() de JavaScript para llamar a agregar_producto_ajax() y actualizar la columna derecha sin recargar la pagina.")

doc.add_page_break()

# ── 9. AUTENTICACION Y ROLES ──────────────────────────────────────────────────

add_heading(doc, "9. Autenticacion y Control de Acceso", 1, border=True)

add_heading(doc, "9.1 Modelo de Autenticacion", 2)
add_para(doc, "El sistema extiende el modelo de autenticacion nativo de Django:")
add_bullet(doc, "Cada empleado tiene un User de Django con username y password para el login.")
add_bullet(doc, "El modelo Empleado esta vinculado al User mediante una relacion OneToOne.")
add_bullet(doc, "El rol del empleado se accede con la expresion: user.empleado.rol.nombre")

add_heading(doc, "9.2 Roles del Sistema", 2)
add_space(doc)
make_table(doc,
    headers=["Rol", "Descripcion", "Acciones permitidas"],
    rows=[
        ["mozo",   "Personal de salon",
         "Abrir mesas, tomar pedidos, marcar items como entregados."],
        ["cajero", "Personal de caja",
         "Emitir documentos tributarios, gestionar caja y arqueos."],
        ["chef",   "Personal de cocina",
         "Ver monitor de cocina, cambiar estados de items (Pendiente -> Cocina -> Listo)."],
    ],
    col_widths_cm=[3.0, 4.5, 8.7]
)
add_space(doc)
add_para(doc, (
    "Las vistas aplican el decorador @login_required de Django y verifican el rol segun "
    "la logica de negocio de cada funcion."
))

doc.add_page_break()

# ── 10. COMANDOS DE OPERACION ─────────────────────────────────────────────────

add_heading(doc, "10. Comandos de Operacion", 1, border=True)

add_heading(doc, "10.1 Configuracion Inicial y Desarrollo", 2)
add_space(doc)
make_table(doc,
    headers=["Comando", "Descripcion"],
    rows=[
        ["env\\Scripts\\activate",              "Activa el entorno virtual Python (Windows)"],
        ["pip install -r requirements.txt",    "Instala todas las dependencias del proyecto"],
        ["python manage.py makemigrations",    "Genera archivos de migracion a partir de los modelos"],
        ["python manage.py migrate",           "Aplica las migraciones a la base de datos PostgreSQL"],
        ["python manage.py showmigrations",    "Muestra el estado de todas las migraciones"],
        ["python manage.py createsuperuser",   "Crea un usuario administrador para el panel Django"],
        ["python manage.py runserver",         "Inicia el servidor de desarrollo en http://localhost:8000"],
        ["python manage.py shell",             "Abre la shell interactiva de Django para pruebas"],
    ],
    col_widths_cm=[6.0, 10.2],
    code_col=0
)

add_space(doc)
add_heading(doc, "10.2 Panel de Administracion", 2)
add_para(doc, (
    "Django Admin esta disponible en http://localhost:8000/admin/ para usuarios con is_staff=True. "
    "Desde alli se pueden gestionar todos los modelos: empleados, mesas, productos, categorias, "
    "timbrados, documentos tributarios y arqueos."
))

doc.add_page_break()

# ── 11. GLOSARIO ──────────────────────────────────────────────────────────────

add_heading(doc, "11. Glosario", 1, border=True)
add_space(doc)

make_table(doc,
    headers=["Termino", "Definicion"],
    rows=[
        ["Comanda",         "Documento o pantalla donde el mozo registra los pedidos de una mesa."],
        ["FIFO",            "First In, First Out. Los pedidos mas antiguos se atienden primero en la cola de cocina."],
        ["Timbrado",        "Numero habilitado por la SET (Paraguay) para emision de comprobantes fiscales."],
        ["Arqueo",          "Conteo y conciliacion del dinero en caja al inicio o cierre del turno."],
        ["FlujoEstado",     "Entidad que registra el estado actual e historial de transiciones de un item de pedido."],
        ["precio_historico","Snapshot del precio de venta al momento de crear el detalle del pedido. Inmutable."],
        ["F-expression",    "Expresion de Django ORM que referencia valores de columna directamente en la BD."],
        ["MVT",             "Model-View-Template. Patron arquitectonico de Django (variante de MVC)."],
        ["ORM",             "Object-Relational Mapper. Capa de abstraccion de Django que mapea clases Python a SQL."],
        ["Bootstrap",       "Framework CSS para diseno responsivo. Version 5.3.0 usada via CDN."],
        ["Japandi",         "Estilo decorativo que fusiona lo japones y escandinavo: minimalismo, tonos neutros."],
    ],
    col_widths_cm=[4.0, 12.2]
)

# ── Guardar ───────────────────────────────────────────────────────────────────

output = r"C:\Proyecto\docs\especificacion_tecnica_ristorantino.docx"
doc.save(output)
print(f"Documento generado: {output}")
